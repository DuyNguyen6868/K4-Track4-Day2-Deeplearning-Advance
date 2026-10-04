"""results.py - dựng results.xlsx (7 sheet, GUIDE mục 6.1), kiểm tra nhất quán, và bảng markdown cho báo cáo.

Mọi con số được ĐỌC từ log chạy thật (runs/*/seed*/summary.json, history.csv, inference_table.csv,
latency_table.csv) và TÍNH LẠI từ predictions/ bằng eval.py gốc cho phần chung kết/mốc (RUBRIC mục I).
Không có số nhập tay. Chạy được ở mọi giai đoạn: phần nào chưa có dữ liệu thì sheet tương ứng để trống.

    python results.py --xlsx results.xlsx          (hoặc results.build_results(...) trong notebook)
    python results.py --check                       (kiểm tra exp_id <-> curves <-> predictions <-> xlsx)
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import experiments as X  # noqa: E402
from evalref import ev  # noqa: E402

FINAL, BASE = "F01", "T00"
RARE = ("Chinee Apple", "Snake Weed")
CLASS_NAMES = ev.CLASS_NAMES


def _pm(mean, std, digits=4):
    return ev.fmt(float(mean), float(std), digits)


def _group(pattern: str, ref_csv: str | None):
    try:
        return ev.load_group(pattern, ref_csv, ref_what="test")
    except FileNotFoundError:
        return None


def _val_f1_from_preds(pattern: str) -> dict[int, float]:
    out = {}
    for f in sorted(Path(pattern).parent.glob(Path(pattern).name)):
        p = ev.read_pred(str(f))
        out[p.seed] = ev.compute_metrics(p.y_true, p.y_pred, p.probs)["macro_f1"]
    return out


def sheet_backbones() -> pd.DataFrame:
    d = X.load_decisions()
    rows = []
    for exp_id, key in X.BACKBONE_RUNS.items():
        s = X.load_summary(exp_id, 0)
        if not s:
            continue
        note = []
        if d.get("backbone_primary") == key:
            note.append("được chọn đi tiếp (backbone chính)")
        if d.get("backbone_light") == key:
            note.append("mạng nhẹ tốt nhất (dùng cho ensemble I05b)")
        rows.append({"exp_id": exp_id, "backbone": key, "họ": X.BACKBONE_FAMILY[key],
                     "tag trọng số": s["weight_tag"], "#tham số (M)": s["params_M"], "GMAC": s["gmacs"],
                     "độ phân giải": s["img_size"], "epoch": s["epochs"], "seed": s["seed"],
                     "macro-F1 val": s["val_macro_f1"], "top-1 val": s["val_top1"],
                     "epoch tốt nhất": s["best_epoch"], "thời gian train/epoch (s)": s["train_time_per_epoch_s"],
                     "độ trễ batch-1 p50 (ms)": s["latency_b1_ms"], "độ trễ batch-1 p95 (ms)": s["latency_b1_p95_ms"],
                     "GPU": s["gpu"], "ghi chú": "; ".join(note)})
    return pd.DataFrame(rows)


def sheet_training() -> pd.DataFrame:
    d = X.load_decisions()
    base0 = X.load_summary(BASE, 0)
    std = d.get("noise", {}).get("std")
    rows = []

    def add(exp_id, axis, changes, s, note=""):
        cls = dict(zip(CLASS_NAMES, s["val_f1_per_class"]))
        delta = s["val_macro_f1"] - base0["val_macro_f1"] if base0 else None
        rows.append({"exp_id": exp_id, "backbone": s["backbone"], "trục thay đổi (A-G)": axis,
                     "khác T00 ở điểm nào": changes, "seed": s["seed"], "macro-F1 val": s["val_macro_f1"],
                     "top-1 val": s["val_top1"], "Δ macro-F1 so với T00": delta,
                     "Δ / std(T00)": (delta / std) if (delta is not None and std and std == std) else None,
                     "F1 val Chinee Apple": cls["Chinee Apple"], "F1 val Snake Weed": cls["Snake Weed"],
                     "F1 val Parkinsonia": cls["Parkinsonia"],
                     "thắng (Δ >= ngưỡng)": d.get("winners_table", {}).get(exp_id, {}).get("winner"),
                     "ghi chú": note})

    for sd in X.FINAL_SEEDS:
        s = X.load_summary(BASE, sd)
        if s:
            add(BASE, "—", "công thức nền (T00)", s, "mốc; 3 seed để đo nhiễu" if sd == 0 else "seed phụ đo nhiễu")
    for e in X.TRAIN_EXPERIMENTS:
        s = X.load_summary(e["exp_id"], 0)
        if s:
            add(e["exp_id"], e["axis"], e["changes"], s, "" if e["adoptable"] else "chỉ phân tích, không ghép vào công thức cuối")
    s = X.load_summary("T20", 0)
    if s:
        members = d.get("combo", {}).get("members", [])
        add("T20", "kết hợp", "kết hợp: " + " + ".join(members), s, "kiểm tra cộng dồn hay triệt tiêu")
    return pd.DataFrame(rows)


def sheet_inference() -> pd.DataFrame:
    p = Path(X.PATHS["out_dir"]) / "inference" / "inference_table.csv"
    if not p.exists():
        return pd.DataFrame()
    t = pd.read_csv(p)
    rename = {"method": "phương pháp", "model": "mô hình/checkpoint", "K": "K (số view hoặc số mô hình)",
              "macro_f1_val": "macro-F1 val", "top1_val": "top-1 val", "ece_val": "ECE val",
              "nll_val": "NLL val", "p50_ms": "p50 batch-1 (ms)", "p95_ms": "p95 batch-1 (ms)",
              "p99_ms": "p99 batch-1 (ms)", "throughput_b1_img_s": "thông lượng batch-1 (ảnh/s)",
              "cost_vs_I00": "chi phí tương đối so với I00", "note": "ghi chú"}
    return t.rename(columns=rename)


def _final_rows(test_csv: str | None, pred_dir: Path):
    rows, groups = [], {}
    cfg_desc = {FINAL: "cấu hình chung kết (backbone + công thức + suy luận + TS)",
                BASE: "mốc: T00 + I00",
                f"{FINAL}uncal": "chung kết, chưa temperature scaling",
                f"{FINAL}rt": "cấu hình thời gian thực (1 view FP32 của mô hình F01)"}
    pats = {FINAL: f"{FINAL}_seed*_test.csv", BASE: f"{BASE}_seed*_test.csv",
            f"{FINAL}uncal": f"{FINAL}_uncal_seed*_test.csv", f"{FINAL}rt": f"{FINAL}rt_seed*_test.csv"}
    val_f1 = _val_f1_from_preds(str(pred_dir / f"{FINAL}_seed*_val.csv"))
    for tag, pat in pats.items():
        g = _group(str(pred_dir / pat), test_csv)
        if g is None:
            continue
        groups[tag] = g
        for p, m in zip(g.preds, g.metrics):
            if tag == FINAL:
                vf = val_f1.get(p.seed)
            elif tag == BASE:
                s = X.load_summary(BASE, p.seed)
                vf = s["val_macro_f1"] if s else None
            else:
                vf = None
            rows.append({"exp_id": tag, "cấu hình": cfg_desc[tag], "seed": p.seed, "macro-F1 val": vf,
                         "macro-F1 test": m["macro_f1"], "top-1 test": m["top1"],
                         "balanced acc test": m["balanced_acc"], "ECE test": m["ece"],
                         "recall Chinee Apple": m["recall"][CLASS_NAMES.index("Chinee Apple")],
                         "recall Snake Weed": m["recall"][CLASS_NAMES.index("Snake Weed")]})
        s_ = g.summary
        rec_m, rec_s = s_["recall"]
        vals = [r["macro-F1 val"] for r in rows if r["exp_id"] == tag and r["macro-F1 val"] is not None]
        vstd = float(np.std(vals, ddof=1)) if len(vals) > 1 else float("nan")
        rows.append({"exp_id": tag, "cấu hình": cfg_desc[tag], "seed": f"mean ± std ({len(g.preds)} seed)",
                     "macro-F1 val": _pm(np.mean(vals), vstd) if vals else None,
                     "macro-F1 test": _pm(*s_["macro_f1"]), "top-1 test": _pm(*s_["top1"]),
                     "balanced acc test": _pm(*s_["balanced_acc"]), "ECE test": _pm(*s_["ece"]),
                     "recall Chinee Apple": _pm(rec_m[CLASS_NAMES.index("Chinee Apple")], rec_s[CLASS_NAMES.index("Chinee Apple")]),
                     "recall Snake Weed": _pm(rec_m[CLASS_NAMES.index("Snake Weed")], rec_s[CLASS_NAMES.index("Snake Weed")])})
    return pd.DataFrame(rows), groups


def sheet_final_perclass(test_csv: str | None, pred_dir: Path):
    final_df, groups = _final_rows(test_csv, pred_dir)
    pcs = []
    names = ev.load_names(str(Path(X.PATHS["labels_dir"]) / "labels.csv")) if (Path(X.PATHS["labels_dir"]) / "labels.csv").exists() else CLASS_NAMES
    for tag in (FINAL, BASE):
        g = groups.get(tag)
        if g is None:
            continue
        for i, n in enumerate(names):
            row = {"cấu hình": tag, "lớp": n, "số ảnh test": int(g.metrics[0]["support"][i])}
            for key, lab in (("precision", "precision"), ("recall", "recall"), ("f1", "F1")):
                row[lab] = _pm(g.summary[key][0][i], g.summary[key][1][i], 3)
            pcs.append(row)
    return final_df, pd.DataFrame(pcs), groups


def sheet_latency() -> pd.DataFrame:
    p = Path(X.PATHS["out_dir"]) / "inference" / "latency_table.csv"
    if not p.exists():
        return pd.DataFrame()
    t = pd.read_csv(p)
    keep = {"config": "cấu hình", "gpu": "GPU", "dtype": "dtype", "batch": "batch", "fused_bn": "gộp BN",
            "img_size": "độ phân giải", "p50": "p50 (ms)", "p95": "p95 (ms)", "p99": "p99 (ms)",
            "images_per_s": "ảnh/s", "warmup": "warmup", "n": "số lần đo", "torch": "torch"}
    t = t[[c for c in keep if c in t.columns]].rename(columns=keep)
    if "gộp BN" in t:
        t["gộp BN"] = t["gộp BN"].map({True: "có", False: "không"}).fillna("không")
    return t


def sheet_summary(back, train, inf, final_df) -> tuple[pd.DataFrame, int | None, int | None]:
    """Top 10 theo macro-F1 val trên mọi thí nghiệm có số val; đánh dấu mốc. Trả về (df, hàng tốt nhất, hàng mốc)."""
    rows = []
    for _, r in back.iterrows():
        rows.append({"exp_id": r["exp_id"], "loại": "backbone", "mô tả": r["backbone"],
                     "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"],
                     "độ trễ p50 batch-1 (ms)": r["độ trễ batch-1 p50 (ms)"], "chi phí tương đối": None})
    for _, r in train.iterrows():
        if r["seed"] != 0:
            continue
        rows.append({"exp_id": r["exp_id"], "loại": "công thức huấn luyện", "mô tả": r["khác T00 ở điểm nào"],
                     "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"],
                     "độ trễ p50 batch-1 (ms)": None, "chi phí tương đối": None})
    for _, r in inf.iterrows() if len(inf) else []:
        if pd.notna(r.get("macro-F1 val")):
            rows.append({"exp_id": r["exp_id"], "loại": "suy luận", "mô tả": r["phương pháp"],
                         "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"],
                         "độ trễ p50 batch-1 (ms)": r.get("p50 batch-1 (ms)"),
                         "chi phí tương đối": r.get("chi phí tương đối so với I00")})
    df = pd.DataFrame(rows)
    if df.empty:
        return df, None, None
    df = df.sort_values("macro-F1 val", ascending=False).head(10).reset_index(drop=True)
    df.insert(0, "hạng", range(1, len(df) + 1))
    # macro-F1 test chỉ có cho chung kết/mốc
    df["macro-F1 test (mean ± std)"] = ""
    if len(final_df):
        for tag in (FINAL, BASE):
            m = final_df[(final_df["exp_id"] == tag) & final_df["seed"].astype(str).str.startswith("mean")]
            if len(m):
                df.loc[df["exp_id"] == tag, "macro-F1 test (mean ± std)"] = m.iloc[0]["macro-F1 test"]
    best = 0
    base = df.index[df["exp_id"] == BASE].tolist()
    return df, best, (base[0] if base else None)


def comparison_block(groups) -> pd.DataFrame:
    """Chung kết so với mốc: Δ macro-F1 test, s, và đạt I2 hay không (cùng định nghĩa với eval.py grade)."""
    if FINAL not in groups or BASE not in groups:
        return pd.DataFrame()
    f, b = groups[FINAL], groups[BASE]
    mf, sf = f.summary["macro_f1"]
    mb, sb = b.summary["macro_f1"]
    s = max([v for v in (sf, sb) if v == v], default=float("nan"))
    delta = mf - mb
    return pd.DataFrame([{"chỉ số": "macro-F1 test", "chung kết (mean)": mf, "mốc T00+I00 (mean)": mb, "Δ": delta,
                          "s = std lớn hơn": s, "Δ > s": bool(delta > s),
                          "điểm I2 đề xuất": ev.points_i2(delta, s)}])


def _style(path: Path, best_rows: dict[str, int], baseline_rows: dict[str, int]) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = load_workbook(path)
    head = PatternFill("solid", fgColor="DDEBF7")
    green = PatternFill("solid", fgColor="C6EFCE")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for ws in wb.worksheets:
        ws.freeze_panes = "B2"
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for col in ws.columns:
            width = max((len(str(c.value)) if c.value is not None else 0) for c in col[:60])
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
            for c in col[1:]:
                if isinstance(c.value, float):
                    c.number_format = "0.0000"
        if ws.title in best_rows:
            for c in ws[best_rows[ws.title] + 2]:
                c.fill = green
        if ws.title in baseline_rows:
            for c in ws[baseline_rows[ws.title] + 2]:
                c.fill = yellow
    wb.save(path)


def build_results(xlsx_path: str = "results.xlsx") -> Path:
    """Dựng results.xlsx từ log chạy thật và predictions/. Trả về đường dẫn file."""
    pred_dir = Path(X.PATHS["pred_dir"])
    test_csv = Path(X.PATHS["labels_dir"]) / "test_subset0.csv"
    test_csv = str(test_csv) if test_csv.exists() else None
    back, train, inf, lat = sheet_backbones(), sheet_training(), sheet_inference(), sheet_latency()
    final_df, perclass, groups = sheet_final_perclass(test_csv, pred_dir)
    summary, best_i, base_i = sheet_summary(back, train, inf, final_df)
    cmp_df = comparison_block(groups)

    path = Path(xlsx_path)
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        for name, df in (("Backbones", back), ("Training", train), ("Inference", inf), ("Final", final_df),
                         ("PerClass", perclass), ("Latency", lat)):
            (df if len(df) else pd.DataFrame({"(chưa có dữ liệu)": []})).to_excel(w, sheet_name=name, index=False)
        (summary if len(summary) else pd.DataFrame({"(chưa có dữ liệu)": []})).to_excel(w, sheet_name="Summary", index=False)
        if len(cmp_df):
            cmp_df.to_excel(w, sheet_name="Summary", index=False, startrow=len(summary) + 3)
    best = {}
    if len(back):
        best["Backbones"] = int(back["macro-F1 val"].idxmax())
    if len(train):
        best["Training"] = int(train["macro-F1 val"].idxmax())
    if len(inf) and inf["macro-F1 val"].notna().any():
        best["Inference"] = int(inf["macro-F1 val"].astype(float).idxmax())
    if best_i is not None:
        best["Summary"] = best_i
    base_rows = {"Summary": base_i} if base_i is not None else {}
    if len(train):
        base_rows["Training"] = int(train.index[(train["exp_id"] == BASE) & (train["seed"] == 0)][0])
    if len(final_df):
        f_idx = final_df.index[(final_df["exp_id"] == FINAL) & final_df["seed"].astype(str).str.startswith("mean")]
        if len(f_idx):
            best["Final"] = int(f_idx[0])
        b_idx = final_df.index[(final_df["exp_id"] == BASE) & final_df["seed"].astype(str).str.startswith("mean")]
        if len(b_idx):
            base_rows["Final"] = int(b_idx[0])
    _style(path, best, base_rows)
    return path


def check_consistency(xlsx_path: str = "results.xlsx") -> list[str]:
    """Trả về danh sách vấn đề (rỗng = ổn): thiếu ảnh curves, thiếu predictions, số trong xlsx lệch số tính lại."""
    problems = []
    curves = Path(X.PATHS["curves_dir"])
    names = [p.name for p in curves.glob("*.png")]
    with pd.ExcelFile(xlsx_path) as xl:
        sheet_names = list(xl.sheet_names)
        sheets = {s: xl.parse(s) for s in ("Backbones", "Training", "Final") if s in sheet_names}
    for sheet in ("Backbones", "Training"):
        if sheet in sheets and "exp_id" in sheets[sheet]:
            for e in sheets[sheet]["exp_id"].dropna().unique():
                if not any(n.startswith(f"{e}_") for n in names):
                    problems.append(f"{sheet}: thiếu ảnh curves/{e}_*.png")
    pred = Path(X.PATHS["pred_dir"])
    for tag in (FINAL, BASE):
        files = list(pred.glob(f"{tag}_seed*_test.csv"))
        if len(files) < 3:
            problems.append(f"predictions: {tag} chỉ có {len(files)} seed (cần >= 3)")
    for e in (FINAL, BASE):
        if not any(n.startswith(f"{e}_") for n in names):
            problems.append(f"thiếu ảnh curves/{e}_*.png")
    # xlsx khớp eval.py
    test_csv = Path(X.PATHS["labels_dir"]) / "test_subset0.csv"
    if "Final" in sheets and test_csv.exists():
        fin = sheets["Final"]
        for tag in (FINAL, BASE):
            g = _group(str(pred / f"{tag}_seed*_test.csv"), str(test_csv))
            if g is None:
                continue
            for p, m in zip(g.preds, g.metrics):
                row = fin[(fin["exp_id"] == tag) & (fin["seed"].astype(str) == str(p.seed))]
                if row.empty:
                    problems.append(f"Final: thiếu dòng {tag} seed {p.seed}")
                elif abs(float(row.iloc[0]["macro-F1 test"]) - m["macro_f1"]) > 1e-9:
                    problems.append(f"Final: macro-F1 test {tag} seed {p.seed} lệch số tính lại từ predictions")
    return problems


def md_table(df: pd.DataFrame, digits: int = 4) -> str:
    if df is None or not len(df):
        return "_(chưa có dữ liệu)_\n"
    def f(v):
        if isinstance(v, (float, np.floating)):
            return "" if v != v else f"{v:.{digits}f}"
        return "" if v is None else str(v)
    head = "| " + " | ".join(map(str, df.columns)) + " |\n|" + "---|" * len(df.columns) + "\n"
    return head + "".join("| " + " | ".join(f(v) for v in r) + " |\n" for r in df.itertuples(index=False))


def write_report_tables(path: str = "report_tables.md") -> Path:
    """Các bảng markdown (sinh tự động từ số thật) để dán vào report.md."""
    pred = Path(X.PATHS["pred_dir"])
    test_csv = Path(X.PATHS["labels_dir"]) / "test_subset0.csv"
    final_df, perclass, groups = sheet_final_perclass(str(test_csv) if test_csv.exists() else None, pred)
    back, train, inf = sheet_backbones(), sheet_training(), sheet_inference()
    cols_b = ["exp_id", "backbone", "tag trọng số", "#tham số (M)", "GMAC", "macro-F1 val", "top-1 val",
              "thời gian train/epoch (s)", "độ trễ batch-1 p50 (ms)", "ghi chú"]
    cols_t = ["exp_id", "trục thay đổi (A-G)", "khác T00 ở điểm nào", "seed", "macro-F1 val", "Δ macro-F1 so với T00",
              "Δ / std(T00)", "F1 val Chinee Apple", "F1 val Snake Weed", "thắng (Δ >= ngưỡng)"]
    cols_i = ["exp_id", "phương pháp", "K (số view hoặc số mô hình)", "macro-F1 val", "top-1 val", "ECE val",
              "p50 batch-1 (ms)", "chi phí tương đối so với I00"]
    parts = ["## Backbone (B)\n", md_table(back[[c for c in cols_b if c in back]] if len(back) else back),
             "\n## Công thức huấn luyện (T)\n", md_table(train[[c for c in cols_t if c in train]] if len(train) else train),
             "\n## Suy luận (I)\n", md_table(inf[[c for c in cols_i if c in inf]] if len(inf) else inf),
             "\n## Chung kết và mốc (test)\n", md_table(final_df),
             "\n## So với mốc\n", md_table(comparison_block(groups)),
             "\n## Theo lớp (test)\n", md_table(perclass)]
    out = Path(path)
    out.write_text("# Bảng sinh tự động từ số thật (results.py)\n\n" + "".join(parts), encoding="utf-8")
    return out


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser()
    for n in X.PATHS:
        ap.add_argument("--" + n.replace("_", "-"))
    ap.add_argument("--xlsx", default="results.xlsx")
    ap.add_argument("--check", action="store_true")
    a = ap.parse_args()
    X.set_paths(**{k: getattr(a, k) for k in X.PATHS})
    if a.check:
        probs = check_consistency(a.xlsx)
        print("\n".join(probs) if probs else "OK: không có vấn đề nhất quán")
        sys.exit(1 if probs else 0)
    print("Đã ghi", build_results(a.xlsx))
    print("Đã ghi", write_report_tables())
